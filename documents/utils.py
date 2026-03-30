import os
import csv
import io
from django.http import HttpResponse
import pdfplumber


# ─── TEXT EXTRACTION ──────────────────────────────────────────────────────────

def extract_text(file, filename):
    """Extract plain text from any uploaded file format."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == '.pdf':
        return extract_from_pdf(file)
    elif ext == '.docx':
        return extract_from_docx(file)
    elif ext in ('.xlsx', '.xls'):
        return extract_from_excel(file)
    elif ext == '.txt':
        return file.read().decode('utf-8', errors='ignore')
    elif ext == '.csv':
        return extract_from_csv(file)
    else:
        # Try reading as plain text for unknown formats
        try:
            return file.read().decode('utf-8', errors='ignore')
        except Exception:
            return ''


def extract_from_pdf(file):
    text = ''
    try:
        with pdfplumber.open(file) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + '\n'
        if not text.strip():
            return "No text could be extracted from this PDF."
        return text.strip()
    except Exception as e:
        print(f"PDF extraction error: {e}")  # logs the error
        return ''

def extract_from_docx(file):
    try:
        from docx import Document as DocxDoc
        doc = DocxDoc(file)
        return '\n'.join([para.text for para in doc.paragraphs if para.text])
    except Exception as e:
        return f'DOCX extraction error: {str(e)}'


def extract_from_excel(file):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file, data_only=True)
        text = ''
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = ' '.join([str(c) for c in row if c is not None])
                if row_text.strip():
                    text += row_text + '\n'
        return text.strip()
    except Exception as e:
        return f'Excel extraction error: {str(e)}'


def extract_from_csv(file):
    try:
        content = file.read().decode('utf-8', errors='ignore')
        reader = csv.reader(io.StringIO(content))
        return '\n'.join([' '.join(row) for row in reader])
    except Exception as e:
        return f'CSV extraction error: {str(e)}'


# ─── EXPORT / DOWNLOAD ────────────────────────────────────────────────────────

def export_document(document, fmt):
    """Generate a downloadable response for the given document in any format."""
    fmt = fmt.lower()

    if fmt == 'pdf':
        return export_as_pdf(document)
    elif fmt == 'docx':
        return export_as_docx(document)
    elif fmt == 'xlsx':
        return export_as_excel(document)
    elif fmt == 'csv':
        return export_as_csv(document)
    else:
        return export_as_txt(document)


def export_as_txt(document):
    response = HttpResponse(content_type='text/plain')
    response['Content-Disposition'] = f'attachment; filename="{document.title}.txt"'
    response.write(f"{document.title}\n\n{document.content}")
    return response


def export_as_csv(document):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{document.title}.csv"'
    writer = csv.writer(response)
    writer.writerow(['Title', 'Type', 'Uploaded By', 'Uploaded At', 'Content'])
    writer.writerow([
        document.title,
        document.doc_type,
        document.uploaded_by,
        document.uploaded_at,
        document.content
    ])
    return response


def export_as_pdf(document):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4,
                                 rightMargin=2*cm, leftMargin=2*cm,
                                 topMargin=2*cm, bottomMargin=2*cm)
        styles = getSampleStyleSheet()
        story = []

        story.append(Paragraph(document.title, styles['Title']))
        story.append(Spacer(1, 0.5*cm))
        story.append(Paragraph(f"Type: {document.doc_type} | By: {document.uploaded_by}", styles['Normal']))
        story.append(Spacer(1, 0.5*cm))

        # Split content into paragraphs
        for para in document.content.split('\n'):
            if para.strip():
                story.append(Paragraph(para.strip(), styles['Normal']))
                story.append(Spacer(1, 0.2*cm))

        doc.build(story)
        buffer.seek(0)

        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{document.title}.pdf"'
        return response
    except Exception as e:
        return HttpResponse(f'PDF export error: {str(e)}', status=500)


def export_as_docx(document):
    try:
        from docx import Document as DocxDoc
        from docx.shared import Pt

        buffer = io.BytesIO()
        doc = DocxDoc()

        doc.add_heading(document.title, 0)
        doc.add_paragraph(f"Type: {document.doc_type} | Uploaded by: {document.uploaded_by}")
        doc.add_paragraph('')

        for para in document.content.split('\n'):
            if para.strip():
                doc.add_paragraph(para.strip())

        doc.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="{document.title}.docx"'
        return response
    except Exception as e:
        return HttpResponse(f'DOCX export error: {str(e)}', status=500)


def export_as_excel(document):
    try:
        import openpyxl
        from openpyxl.styles import Font

        buffer = io.BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Document'

        # Header row
        ws['A1'] = 'Title'
        ws['B1'] = 'Type'
        ws['C1'] = 'Uploaded By'
        ws['D1'] = 'Uploaded At'
        ws['A1'].font = Font(bold=True)
        ws['B1'].font = Font(bold=True)
        ws['C1'].font = Font(bold=True)
        ws['D1'].font = Font(bold=True)

        # Data row
        ws['A2'] = document.title
        ws['B2'] = document.doc_type
        ws['C2'] = document.uploaded_by
        ws['D2'] = str(document.uploaded_at)

        # Content sheet
        ws2 = wb.create_sheet('Content')
        ws2['A1'] = 'Content'
        ws2['A1'].font = Font(bold=True)
        for i, line in enumerate(document.content.split('\n'), start=2):
            ws2.cell(row=i, column=1, value=line)

        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{document.title}.xlsx"'
        return response
    except Exception as e:
        return HttpResponse(f'Excel export error: {str(e)}', status=500)